from io import BytesIO
import hashlib
import unittest
import fitz
from openpyxl import load_workbook
from template_fill_test_v2 import Estimate, Category, LineItem, fill_estimate_v2, TEMPLATE_PATH
from pour_layout import plan, describe


def example(counts=(12,8,12), sources=2, expense=False):
    trades=('防水工事','足場工事','塗装工事')
    categories=[]
    for index,count in enumerate(counts):
        categories.append(Category(trades[index],[LineItem(f'{trades[index]}作業{i+1}','確認仕様',10,'㎡',1000,10000,row_id=f'r{index}-{i}') for i in range(count)]))
    if expense:
        categories.append(Category('原本諸経費',[LineItem('原本運搬費','原本費目',1,'式',1000,1000,row_id='expense')],True))
    return Estimate('非公開の施工業者名','架空確認先 御中','配置確認（架空）','2026/10/20','30日',
        categories,-1000,.1,pour_layout=True,source_count=sources)


class PourLayoutTest(unittest.TestCase):
    def output(self,e):
        result=fill_estimate_v2(e,save_to_disk=False)
        self.assertTrue(result.ok,result.errors)
        return result,load_workbook(BytesIO(result.data))

    def test_main_trades_start_on_separate_pages(self):
        e=example(); layout=plan(e)
        self.assertFalse(layout['compact'])
        self.assertEqual(len(describe(e)),5)
        for page,trade in zip(layout['pages'],('防水工事','足場工事','塗装工事')):
            self.assertTrue(all(r['item'].name.startswith(trade) for r in page if r['kind']=='item'))
        result,w=self.output(e)
        self.assertEqual(w.sheetnames,['見積書','工事品目集計','明細1','明細2','明細3'])
        with fitz.open(stream=result.pdf_data,filetype='pdf') as pdf:
            self.assertEqual(len(pdf),5)
            self.assertNotIn('足場工事',pdf[2].get_text())
            self.assertIn('足場工事',pdf[3].get_text())

    def test_two_row_trade_shares_detail_but_keeps_separate_summary(self):
        e=example((12,2)); result,w=self.output(e)
        self.assertEqual(len(plan(e)['pages']),1)
        self.assertEqual(w['工事品目集計']['B3'].value,'防水工事')
        self.assertEqual(w['工事品目集計']['B4'].value,'足場工事')
        self.assertNotEqual(w['工事品目集計']['G3'].value,w['工事品目集計']['G4'].value)

    def test_tiny_trade_does_not_become_a_two_row_last_page(self):
        e=example((18,2)); pages=plan(e)['pages']
        self.assertEqual(len(pages),2)
        self.assertGreater(sum(r['kind']=='item' for r in pages[-1]),2)
        self.assertEqual([r['no'] for p in pages for r in p if r['kind']=='item'],list(range(1,21)))
        self.assertTrue(all(len(p)<=20 for p in pages))

    def test_leading_tiny_trade_can_share_a_large_trades_first_page(self):
        pages=plan(example((2,25)))['pages']
        self.assertGreater(sum(r['kind']=='item' for r in pages[0]),2)
        self.assertTrue(all(len(p)<=20 for p in pages))
        self.assertEqual([r['no'] for p in pages for r in p if r['kind']=='item'],list(range(1,28)))

    def test_single_source_small_estimate_has_cover_and_detail_with_expenses(self):
        e=example((12,),sources=1,expense=True); result,w=self.output(e)
        self.assertTrue(plan(e)['compact'])
        self.assertEqual(w.sheetnames,['見積書','明細1'])
        self.assertEqual(w['見積書']['C13'].value,"='明細1'!G24")
        self.assertEqual(w['明細1']['B15'].value,'原本運搬費')
        with fitz.open(stream=result.pdf_data,filetype='pdf') as pdf:
            self.assertEqual(len(pdf),2)
            self.assertIn('原本運搬費',pdf[1].get_text())
            self.assertIn('税込合計',pdf[1].get_text())

    def test_large_single_source_retains_summary_and_multiple_detail_pages(self):
        e=example((25,),sources=1,expense=True); result,w=self.output(e)
        self.assertFalse(plan(e)['compact'])
        self.assertEqual(len(w.sheetnames),4)
        self.assertEqual(w['工事品目集計']['B4'].value,'原本運搬費')

    def test_multiple_sources_keep_summary_even_with_one_trade(self):
        e=example((1,),sources=2)
        self.assertFalse(plan(e)['compact'])
        self.assertEqual(len(describe(e)),3)

    def test_user_break_can_put_no18_on_page4_without_changing_amounts(self):
        e=example(); total=e.grand_total
        e.page_breaks=('r1-5',);e.requested_pages={'r1-5':4}
        layout=plan(e)
        self.assertEqual(len(layout['pages']),2)
        self.assertEqual(layout['pages'][1][1]['item'].row_id,'r1-5')
        self.assertEqual(e.grand_total,total)
        self.output(e)

    def test_impossible_requested_page_is_rejected_instead_of_compressing(self):
        e=example();e.page_breaks=('r2-0',);e.requested_pages={'r2-0':4}
        with self.assertRaisesRegex(ValueError,'指定'):
            plan(e)
        result=fill_estimate_v2(e,save_to_disk=False)
        self.assertFalse(result.ok)
        self.assertIsNone(result.data)

    def test_template_is_unchanged_and_supplier_name_does_not_leave_app(self):
        before=hashlib.sha256(TEMPLATE_PATH.read_bytes()).hexdigest()
        result,w=self.output(example(expense=True))
        self.assertEqual(hashlib.sha256(TEMPLATE_PATH.read_bytes()).hexdigest(),before)
        values=' '.join(str(c.value or '') for s in w for row in s for c in row)
        self.assertNotIn('非公開の施工業者名',values)
        with fitz.open(stream=result.pdf_data,filetype='pdf') as pdf:
            self.assertNotIn('非公開の施工業者名',''.join(p.get_text() for p in pdf))

    def test_shared_pdf_xlsx_plan_has_no_orphan_heading_or_unprinted_item(self):
        for counts in ((16,2),(17,2),(19,1),(2,25),(12,8,12),(25,3,2)):
            e=example(counts)
            pages=plan(e)['pages']
            with self.subTest(counts=counts):
                self.assertTrue(all(len(p)<=20 for p in pages))
                self.assertTrue(all(p[-1]['kind'] not in ('category','section') for p in pages))
                self.assertTrue(all(p[0]['kind']!='subtotal' for p in pages))
                self.assertEqual([r['no'] for p in pages for r in p if r['kind']=='item'],list(range(1,sum(counts)+1)))
                result,w=self.output(e)
                with fitz.open(stream=result.pdf_data,filetype='pdf') as pdf:
                    self.assertEqual(len(pdf),len(w.sheetnames))

    def test_summary_continues_without_losing_the_seventeenth_category(self):
        e=example((1,),sources=2)
        e.categories=[Category(f'工種{n}',[LineItem(f'確認作業{n}','仕様',1,'式',1000,1000,row_id=f'x{n}')]) for n in range(17)]
        result,w=self.output(e)
        self.assertEqual(w.sheetnames[:3],['見積書','工事品目集計','工事品目集計2'])
        self.assertEqual(w['工事品目集計']['B11'].value,'工種8')
        self.assertEqual(w['工事品目集計2']['B3'].value,'工種9')
        self.assertEqual(w['工事品目集計2']['B10'].value,'工種16')
        self.assertIn("'工事品目集計2'!G3:G18",w['工事品目集計']['G22'].value)
        self.assertEqual(plan(e)['first_detail'],4)
        with fitz.open(stream=result.pdf_data,filetype='pdf') as pdf:
            self.assertEqual(len(pdf),len(w.sheetnames))
            self.assertIn('続き',pdf[2].get_text())
            self.assertIn('工事内',pdf[3].get_text().replace(' ',''))

    def test_more_than_eight_detail_pages_copy_the_existing_design(self):
        e=example((180,),sources=2)
        before=hashlib.sha256(TEMPLATE_PATH.read_bytes()).hexdigest()
        result,w=self.output(e)
        self.assertGreater(len(plan(e)['pages']),8)
        self.assertIn('明細9',w.sheetnames)
        numbers=[sheet.cell(row,1).value for sheet in w if sheet.title.startswith('明細') for row in range(3,23) if isinstance(sheet.cell(row,1).value,int)]
        self.assertEqual(numbers,list(range(1,181)))
        self.assertEqual(w['明細9'].page_setup.orientation,'landscape')
        self.assertEqual(hashlib.sha256(TEMPLATE_PATH.read_bytes()).hexdigest(),before)
        with fitz.open(stream=result.pdf_data,filetype='pdf') as pdf:self.assertEqual(len(pdf),len(w.sheetnames))

    def test_fractional_quantity_and_unit_price_are_not_hidden_by_integer_format(self):
        e=example((1,),sources=1)
        item=e.categories[0].items[0]
        item.qty=.125;item.unit_price=2000.5;item.amount=250
        result,w=self.output(e)
        self.assertIn('#',w['明細1']['D3'].number_format.split('.')[1])
        self.assertIn('#',w['明細1']['F3'].number_format.split('.')[1])
        with fitz.open(stream=result.pdf_data,filetype='pdf') as pdf:
            self.assertIn('0.125',pdf[1].get_text())
            self.assertIn('2,000.5',pdf[1].get_text())


if __name__=='__main__':unittest.main()
