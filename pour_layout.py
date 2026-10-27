"""One shared customer page plan, applied to the existing CYCA template."""
from io import BytesIO
from print_text_height import entry_height, values_height, grid_fits


def _fits(rows):
    return grid_fits([entry_height(r) for r in rows])


def _chunk(rows):
    pages = []
    rows = list(rows)
    while rows:
        prefix = []
        if rows[0]['kind'] == 'item':
            from template_fill_test_v2 import detail_context
            item = rows[0]['item']
            category = rows[0].get('category')
            label = (category.name if category else '工事内容') + '（続き）'
            context = detail_context(item, category) if category else item.section
            if context:
                label += '／' + context
            prefix = [{'kind': 'section', 'label': label}]
        take = 0
        while take < len(rows) and _fits(prefix+rows[:take+1]):
            take += 1
        if take == 0:
            raise ValueError('1品目の記載が1ページに収まりません。原本を保全し、仕様を別紙へ分けてください。')
        if take < len(rows) and rows[take]['kind'] == 'subtotal' and take > 1:
            take -= 1
        while take > 1 and rows[take - 1]['kind'] in ('category', 'section'):
            take -= 1
        pages.append(prefix + rows[:take])
        rows = rows[take:]
    return pages


def summary_chunks(categories):
    count=max(1,(len(categories)+15)//16)
    base,extra=divmod(len(categories),count)
    result=[];offset=0
    for number in range(count):
        size=base+(number<extra)
        result.append(categories[offset:offset+size]);offset+=size
    if all(grid_fits([values_height([i+1,c.items[0].name,c.items[0].spec,None,None,None,None,c.items[0].remark])
                      if c.summary_only else values_height([i+1,c.name]) for i,c in enumerate(chunk)],16)
           for chunk in result):return result
    result=[];chunk=[]
    for category in categories:
        item=category.items[0]
        height=(values_height([1,item.name,item.spec,None,None,None,None,item.remark])
                if category.summary_only else values_height([1,category.name]))
        if not grid_fits([height],16):
            raise ValueError('一式まとめの1品目が1ページに収まりません。仕様を別紙へ分けてください。')
        if chunk and not grid_fits([h for _,h in chunk]+[height],16):
            result.append([c for c,_ in chunk]);chunk=[]
        chunk.append((category,height))
    if chunk:result.append([c for c,_ in chunk])
    return result


def plan(estimate):
    from template_fill_test_v2 import _build_detail_rows, _detail_item_rows
    normal = [c for c in estimate.categories if not c.summary_only]
    all_items = ([i for c in normal for i in c.items]
                 + [i for c in estimate.categories if c.summary_only for i in c.items])
    compact_rows=[]
    ordered=normal+[c for c in estimate.categories if c.summary_only]
    number=0
    for index,category in enumerate(ordered,1):
        if len(normal)>1:
            compact_rows.append({'kind':'category','index':index,'category':category})
        item_rows, number = _detail_item_rows(category, number)
        compact_rows.extend(item_rows)
    compact = (estimate.source_count == 1 and bool(normal) and len(compact_rows) + 4 <= 20
               and _fits(compact_rows+[{'kind':'compact_'+n} for n in ('subtotal','discount','net','tax')])
               and not estimate.page_breaks)
    if compact:
        rows = compact_rows
        rows += [{'kind': 'compact_' + name, 'amount': value} for name, value in
                 [('subtotal', estimate.subtotal), ('discount', estimate.discount),
                  ('net', estimate.net_total), ('tax', estimate.tax)]]
        return {'compact': True, 'pages': [rows], 'first_detail': 2, 'summary_pages':0}
    logical = _build_detail_rows(estimate)
    pages = []
    if estimate.page_breaks:
        known = {r['item'].row_id for r in logical if r['kind'] == 'item'}
        if len(set(estimate.page_breaks)) != len(estimate.page_breaks) or any(k not in known or not k for k in estimate.page_breaks):
            raise ValueError('改ページする品目を確認してください。')
        blocks, block = [], []
        for row in logical:
            if row['kind'] == 'item' and row['item'].row_id in estimate.page_breaks and block:
                headings = []
                while block and block[-1]['kind'] in ('category', 'section'):
                    headings.insert(0, block.pop())
                if block:
                    blocks.append(block)
                block = headings
            block.append(row)
        if block:
            blocks.append(block)
        for block in blocks:
            pages.extend(_chunk(block))
    else:
        # Main trades begin on their own page. A tiny trade may share a page,
        # while retaining its own heading, subtotal and summary entry.
        blocks = []
        for row in logical:
            if row['kind'] == 'category':
                blocks.append([])
            blocks[-1].append(row)
        tiny_page = False
        for block in blocks:
            tiny = sum(r['kind'] == 'item' for r in block) <= 2
            if pages and _fits(pages[-1]+block) and (tiny or tiny_page):
                pages[-1].extend(block)
                tiny_page = tiny_page and tiny
            elif pages and tiny_page and len(pages[-1]) <= 8:
                # Let a larger trade share the leading tiny trade's page.
                combined = _chunk(pages.pop() + block)
                pages.extend(combined)
                tiny_page = False
            elif pages and tiny and len(pages[-1]) >= 12:
                # A two-row trade must not become a mostly empty last page.
                previous = pages.pop()
                cut = len(previous) - 6
                while cut > 1 and previous[cut - 1]['kind'] in ('category', 'section'):
                    cut -= 1
                pages.append(previous[:cut])
                pages.extend(_chunk(previous[cut:] + block))
                tiny_page = False
            else:
                pages.extend(_chunk(block))
                tiny_page = tiny
    summary_pages=len(summary_chunks(estimate.categories))
    first_detail=2+summary_pages
    actual = {r['item'].row_id: n for n, page in enumerate(pages, first_detail) for r in page if r['kind'] == 'item'}
    for row_id, number in estimate.requested_pages.items():
        if row_id not in actual or actual[row_id] != number:
            selected = next((r for r in logical if r['kind']=='item' and r['item'].row_id==row_id), None)
            label = f"No.{selected['no']} {selected['item'].name}" if selected else '指定の品目'
            raise ValueError(f'{label}は現在{actual.get(row_id, "不明")}ページ目です。指定した{number}ページ目には合いません。区切りを変更するか、自動のページ分けに戻してください。')
    return {'compact': False, 'pages': pages, 'first_detail': first_detail, 'summary_pages':summary_pages}


def describe(estimate, layout=None):
    layout = layout or plan(estimate)
    result = [{'page': 1, 'label': '表紙・総合金額', 'rows': []}]
    if not layout['compact']:
        for index in range(layout['summary_pages']):
            result.append({'page':2+index,'label':'工種ごとの一式まとめ・諸経費'+('（続き）' if index else ''),'rows':[]})
    for number, page in enumerate(layout['pages'], layout['first_detail']):
        items = [r for r in page if r['kind'] == 'item']
        trades = list(dict.fromkeys(r['category'].name for r in page if r.get('category') is not None))
        if layout['compact']:
            trades = [c.name for c in estimate.categories if not c.summary_only]
            if any(c.summary_only for c in estimate.categories):
                trades.append('諸経費等')
        result.append({'page': number, 'label': '・'.join(trades) or '工事内容明細',
                       'rows': [{'no': r['no'], 'row_id': r['item'].row_id, 'name': r['item'].name} for r in items]})
    return result


def build_outputs(estimate, template_path):
    from openpyxl import load_workbook
    from openpyxl.styles import PatternFill
    from template_fill_test_v2 import (QUOTE_SHEET, SUMMARY_SHEET, DETAIL_PAGE_START_ROW,
        DETAIL_PAGE_BLOCK_ROWS, MAX_DETAIL_PAGES, _set, BRAND_TINT)
    from print_layout import split_print_pages, render_estimate_pdf
    before = template_path.read_bytes()
    layout = plan(estimate)
    actual_items = [r['item'] for page in layout['pages'] for r in page if r['kind'] == 'item']
    expected_items = ([i for c in estimate.categories if not c.summary_only for i in c.items]
                      + ([i for c in estimate.categories if c.summary_only for i in c.items] if layout['compact'] else []))
    if [id(i) for i in actual_items] != [id(i) for i in expected_items]:
        raise ValueError('明細の欠落・重複・並び順の不一致を検出しました。')
    wb = load_workbook(BytesIO(before))
    source = wb[QUOTE_SHEET]
    for n in range(MAX_DETAIL_PAGES):
        start = DETAIL_PAGE_START_ROW + n * DETAIL_PAGE_BLOCK_ROWS
        for row in range(start, start + DETAIL_PAGE_BLOCK_ROWS):
            source.row_dimensions[row].hidden = n >= len(layout['pages'])
    split_print_pages(wb, estimate)
    if layout['compact']:
        del wb[SUMMARY_SHEET]
    else:
        for n in range(2,layout['summary_pages']+1):
            sheet=wb.copy_worksheet(wb[SUMMARY_SHEET]);sheet.title=SUMMARY_SHEET+str(n)
            wb.move_sheet(sheet,offset=(n-len(wb.worksheets)+1))
    for n in range(1,len(layout['pages'])+1):
        if f'明細{n}' not in wb:
            sheet=wb.copy_worksheet(wb['明細1']);sheet.title=f'明細{n}'
    cover = wb[QUOTE_SHEET]
    import re
    honorific = re.search(r'(御中|様|殿)\s*$', estimate.customer)
    _set(cover, 5, 1, re.sub(r'[\s　]*(御中|様|殿)\s*$', '', estimate.customer))
    _set(cover, 5, 4, honorific.group(1) if honorific else '御中')
    _set(cover, 7, 7, estimate.issue_date)
    _set(cover, 25, 3, estimate.project_name)
    _set(cover, 29, 3, estimate.valid_days)
    item_refs = {i: [] for i in range(len(estimate.categories))}
    category_index = {id(c): i for i, c in enumerate(estimate.categories)}
    item_categories = {id(item): i for i, c in enumerate(estimate.categories) for item in c.items}
    subtotal_cells = []
    compact_cells = {}
    for n, page in enumerate(layout['pages'], 1):
        ws = wb[f'明細{n}']
        for row in range(3, 23):
            for col in range(1, 9):
                _set(ws, row, col, None)
        page_refs = []
        for row, entry in enumerate(page, 3):
            kind = entry['kind']
            if kind == 'item':
                item = entry['item']
                values = [entry['no'], item.name, item.spec, item.qty, item.unit, item.unit_price, None, item.remark]
                for col, value in enumerate(values, 1):
                    _set(ws, row, col, value)
                ws.cell(row, 4).number_format = '#,##0.################'
                ws.cell(row, 6).number_format = '#,##0.################'
                ws.cell(row, 7).value = f'=ROUND(D{row}*F{row},0)'
                ref = f"'{ws.title}'!G{row}"
                item_refs[item_categories[id(item)]].append(ref)
                page_refs.append(f'G{row}')
            elif kind in ('category', 'section'):
                label = f"{entry['index']}. {entry['category'].name}" if kind == 'category' else '【' + entry['label'] + '】'
                _set(ws, row, 2, label)
                for col in range(1, 9):
                    ws.cell(row, col).fill = PatternFill('solid', fgColor=BRAND_TINT)
            elif kind == 'subtotal':
                _set(ws, row, 6, '小計')
                subtotal_cells.append((ws, row, category_index[id(entry['category'])]))
            else:
                label = {'compact_subtotal': '小計', 'compact_discount': '値引き', 'compact_net': '税抜合計', 'compact_tax': '消費税'}[kind]
                _set(ws, row, 6, label)
                compact_cells[kind] = row
        if layout['compact']:
            subtotal, discount, net, tax = (compact_cells['compact_' + key] for key in ('subtotal', 'discount', 'net', 'tax'))
            ws.cell(subtotal, 7).value = '=SUM(' + ','.join(page_refs) + ')'
            _set(ws, discount, 7, estimate.discount)
            ws.cell(net, 7).value = f'=G{subtotal}+G{discount}'
            ws.cell(tax, 7).value = f'=ROUNDDOWN(G{net}*{estimate.tax_rate},0)'
            _set(ws, 24, 6, '税込合計')
            _set(ws, 23, 1, '各明細金額は税抜です。消費税は合計欄に計上しています。')
            ws.cell(24, 7).value = f'=G{net}+G{tax}'
        else:
            ws.cell(24, 7).value = '=SUM(' + ','.join(page_refs) + ')' if page_refs else '=0'
    for ws, row, index in subtotal_cells:
        ws.cell(row, 7).value = '=SUM(' + ','.join(item_refs[index]) + ')'
    if not layout['compact']:
        summary_names=[SUMMARY_SHEET+(str(n) if n>1 else '') for n in range(1,layout['summary_pages']+1)]
        chunks=summary_chunks(estimate.categories)
        category_offset=0
        for page_index,name in enumerate(summary_names):
            summary=wb[name]
            if page_index:_set(summary,1,1,'工事品目まとめ（続き）')
            for row in range(3,19):
                for col in range(1,9):_set(summary,row,col,None)
            for offset,category in enumerate(chunks[page_index]):
                index=category_offset+offset;row=3+offset
                if category.summary_only:
                    item=category.items[0]
                    values=[index+1,item.name,item.spec,item.qty,item.unit,item.unit_price,None,item.remark]
                    formula=f'=ROUND(D{row}*F{row},0)'
                else:
                    values=[index+1,category.name,'工事一式',1,'式',None,None,'']
                    formula='=SUM('+','.join(item_refs[index])+')'
                for col,value in enumerate(values,1):_set(summary,row,col,value)
                summary.cell(row,4).number_format='#,##0.################'
                summary.cell(row,6).number_format='#,##0.################'
                summary.cell(row,7).value=formula
            category_offset+=len(chunks[page_index])
            if page_index==0:
                summary['G22']='=SUM('+','.join("'"+title+"'!G3:G18" for title in summary_names)+')'
                _set(summary,23,7,estimate.discount)
                summary['G24']='=G22+G23'
                summary['G25']=f'=ROUNDDOWN(G24*{estimate.tax_rate},0)'
                summary['G26']='=G24+G25'
            else:
                for row in (22,23,24,25,26):summary.cell(row,7).value=f"='{SUMMARY_SHEET}'!G{row}"
            if len(summary_names)>1:
                _set(summary,22,6,'全体小計')
                _set(summary,21,1,'合計欄は、まとめ全ページを合わせた金額です。')
    cover['C13'] = "='明細1'!G24" if layout['compact'] else "='工事品目集計'!G26"
    from print_layout import fit_numbers_text_rows
    for n, ws in enumerate(wb, 1):
        ws.oddFooter.right.text = f'{n} / {len(wb.worksheets)}'
        if ws.title.startswith(SUMMARY_SHEET):
            fit_numbers_text_rows(ws, 3, 18)
        elif ws.title.startswith('明細'):
            fit_numbers_text_rows(ws, 3, 22)
    wb.calculation.calcMode = 'auto'
    wb.calculation.fullCalcOnLoad = True
    output = BytesIO()
    wb.save(output)
    if template_path.read_bytes() != before:
        raise ValueError('元の流し込みテンプレートが変更されています。')
    return output.getvalue(), render_estimate_pdf(estimate), wb.sheetnames
