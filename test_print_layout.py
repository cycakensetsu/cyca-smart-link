import unittest
from io import BytesIO
from zipfile import ZipFile
import xml.etree.ElementTree as ET
import fitz
from openpyxl import load_workbook
from template_fill_test_v2 import Estimate, Category, LineItem, fill_estimate_v2, COMPANY

class PrintLayoutTest(unittest.TestCase):
    def estimate(self, count=14):
        return Estimate(vendor='検証工務店', customer='塚本 様', project_name='今回の工事', issue_date='2026/10/05', valid_days='30日', categories=[Category('内装工事', [LineItem(f'工事{i}', '', 2, '㎡', 5000, 10000, '') for i in range(count)])])

    def test_numbers_project_footer_keeps_text_without_size_code_and_preserves_right_pager(self):
        for pour, sources, count in ((False, 2, 25), (True, 2, 25), (True, 1, 3)):
            with self.subTest(pour_layout=pour, source_count=sources):
                estimate=self.estimate(count)
                estimate.project_name='5件工事（確認用）'
                estimate.pour_layout=pour
                estimate.source_count=sources
                result=fill_estimate_v2(estimate,save_to_disk=False)
                self.assertTrue(result.ok,result.errors)
                workbook=load_workbook(BytesIO(result.data))
                with ZipFile(BytesIO(result.data)) as archive:
                    for number,sheet in enumerate(workbook,1):
                        expected_footer=f"{COMPANY['name']}　{estimate.project_name}"
                        self.assertEqual(sheet.oddFooter.left.text,expected_footer)
                        self.assertIsNone(sheet.oddFooter.left.size)
                        self.assertEqual(sheet.oddFooter.right.text,f'{number} / {len(workbook.worksheets)}')
                        self.assertEqual(sheet.oddFooter.right.size,8)
                        root=ET.fromstring(archive.read(f'xl/worksheets/sheet{number}.xml'))
                        footer=root.find('{*}headerFooter/{*}oddFooter').text
                        self.assertEqual(footer.split('&R',1)[0],'&L'+expected_footer)

    def test_print_pages_and_honorific_match_pdf(self):
        result=fill_estimate_v2(self.estimate(), save_to_disk=False)
        self.assertTrue(result.ok,result.errors)
        wb=load_workbook(BytesIO(result.data))
        self.assertEqual(wb.sheetnames,['見積書','工事品目集計','明細1'])
        self.assertEqual(wb['見積書']['D5'].value,'様')
        self.assertEqual(wb['見積書']['C13'].value,"='工事品目集計'!G26")
        for sheet in wb:
            self.assertEqual(sheet.sheet_state,'visible')
            self.assertEqual(sheet.page_setup.fitToWidth,1)
            self.assertEqual(sheet.page_setup.fitToHeight,1)
            self.assertEqual(sheet.page_setup.orientation,'landscape')
            self.assertTrue(sheet.print_area)
        pdf=fitz.open(stream=result.pdf_data,filetype='pdf')
        self.assertEqual(len(pdf),3)
        for page in pdf:
            self.assertAlmostEqual(page.rect.width,841.89,places=1)
            self.assertAlmostEqual(page.rect.height,595.28,places=1)
        self.assertIn('塚本',pdf[0].get_text())
        self.assertNotIn('工 事 内 容 明 細',pdf[1].get_text())
        self.assertIn('工事13',pdf[2].get_text())

    def test_multi_page_subtotal_keeps_cross_sheet_references(self):
        result=fill_estimate_v2(self.estimate(81),save_to_disk=False)
        self.assertTrue(result.ok,result.errors)
        wb=load_workbook(BytesIO(result.data))
        self.assertEqual(wb.sheetnames,['見積書','工事品目集計']+[f'明細{i}' for i in range(1,6)])
        self.assertIn("'明細5'!",wb['工事品目集計']['G3'].value)
        formulas=[c.value for sheet in wb for row in sheet for c in row if c.data_type=='f']
        self.assertFalse(any('_layout_source' in f or '明細データ' in f for f in formulas))
        pdf=fitz.open(stream=result.pdf_data,filetype='pdf')
        self.assertEqual(len(pdf),7)
        self.assertIn('工事80',pdf[-1].get_text())
        self.assertNotIn('工事80',pdf[-2].get_text())

    def test_original_design_keeps_template_and_twenty_row_page_rule(self):
        from template_fill_test_v2 import TEMPLATE_PATH, _build_detail_rows, _paginate_detail_rows, DETAIL_PAGE_DATA_ROWS
        before=TEMPLATE_PATH.read_bytes()
        e=self.estimate(35)
        e.categories.append(Category('諸経費',[LineItem('原本運搬費','原本の費目',1,'式',1000,1000)],True))
        pages=_paginate_detail_rows(_build_detail_rows(e))
        self.assertEqual(DETAIL_PAGE_DATA_ROWS,20)
        self.assertTrue(all(len(p)<=20 for p in pages))
        self.assertTrue(all(p[-1]['kind']!='category' for p in pages))
        self.assertTrue(all(p[0]['kind']!='subtotal' for p in pages))
        result=fill_estimate_v2(e,save_to_disk=False)
        self.assertTrue(result.ok,result.errors)
        self.assertEqual(TEMPLATE_PATH.read_bytes(),before)
        w=load_workbook(BytesIO(result.data))
        self.assertEqual(w['工事品目集計']['G4'].value,'=ROUND(D4*F4,0)')
        self.assertEqual(w.sheetnames,['見積書','工事品目集計','明細1','明細2'])
        with fitz.open(stream=result.pdf_data,filetype='pdf') as doc:
            self.assertEqual(len(doc),4)
            self.assertIn('原本運搬費',doc[1].get_text())
            self.assertNotIn('原本運搬費',doc[2].get_text()+doc[3].get_text())
            self.assertIn('工事34',doc[3].get_text())

    def test_sixteen_summary_categories_stay_inside_page_without_small_fonts(self):
        e=self.estimate(1)
        e.categories=[Category(f'工種{i}',[LineItem('確認用作業','仕様',1,'式',10000,10000)]) for i in range(16)]
        result=fill_estimate_v2(e,save_to_disk=False)
        self.assertTrue(result.ok,result.errors)
        with fitz.open(stream=result.pdf_data,filetype='pdf') as doc:
            blocks=doc[1].get_text('dict')['blocks']
            spans=[s for b in blocks if 'lines' in b for line in b['lines'] for s in line['spans']]
            self.assertTrue(all(s['bbox'][1]>=0 and s['bbox'][3]<=doc[1].rect.height for s in spans))
            self.assertTrue(all(s['size']>=9 for s in spans if any(word in s['text'] for word in ('工種','税込合計','税抜合計'))))

    def test_original_sections_survive_trade_grouping_and_page_breaks(self):
        from template_fill_test_v2 import _build_detail_rows, _paginate_detail_rows
        e=self.estimate(22)
        for i,item in enumerate(e.categories[0].items):item.section='屋上' if i<17 else '塔屋'
        pages=_paginate_detail_rows(_build_detail_rows(e))
        self.assertTrue(all(p[-1]['kind'] not in ('category','section') for p in pages))
        result=fill_estimate_v2(e,save_to_disk=False)
        self.assertTrue(result.ok,result.errors)
        with fitz.open(stream=result.pdf_data,filetype='pdf') as doc:
            detail=''.join(p.get_text() for p in list(doc)[2:])
            self.assertIn('【屋上】',detail)
            self.assertIn('【塔屋】',detail)
            self.assertIn('工事21',detail)

    def test_page_boundary_does_not_leave_a_small_category_heading_alone(self):
        from template_fill_test_v2 import _build_detail_rows, _paginate_detail_rows
        for count in range(16,22):
            e=self.estimate(count)
            e.categories.append(Category('足場工事',[LineItem('作業足場','仕様',1,'式',10000,10000,section='外周')]))
            pages=_paginate_detail_rows(_build_detail_rows(e))
            with self.subTest(count=count):
                self.assertTrue(all(p[-1]['kind'] not in ('category','section') for p in pages))
                self.assertTrue(all(p[0]['kind']!='subtotal' for p in pages))

if __name__=='__main__': unittest.main()
