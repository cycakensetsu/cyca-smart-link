from io import BytesIO
from pathlib import Path
import tempfile
import unittest

import fitz
from openpyxl import load_workbook

from cyca_reference import read_reference_quote
from reference_layout import build_reference_outputs, page_plan, SUMMARY_SHEET
from template_fill_test_v2 import Category, LineItem, Estimate


class ReferenceLayoutTest(unittest.TestCase):
    def estimate(self,count=30):
        items=[LineItem(f'検証工事{i}', '仕様A', .25, '㎡', 4000, 1000) for i in range(count)]
        return Estimate('','確認用 御中','検証用工事','2026/10/20','30日',
                        [Category('防水工事',items),Category('運搬費',[LineItem('運搬費','原本の費目',1,'式',10000,10000)],True)],-1000,.1,True)

    def test_no_blank_detail_page_and_continuous_numbers(self):
        e=self.estimate();xlsx,pdf,names=build_reference_outputs(e,draft=True)
        self.assertEqual(names,['見積書',SUMMARY_SHEET,'明細1','明細2'])
        w=load_workbook(BytesIO(xlsx))
        numbers=[w[s].cell(r,1).value for s in names[2:] for r in range(3,29) if isinstance(w[s].cell(r,1).value,int)]
        self.assertEqual(numbers,list(range(1,31)))
        self.assertEqual(w['見積書']['D16'].value,f"='{SUMMARY_SHEET}'!H31")
        self.assertEqual(w[SUMMARY_SHEET]['H31'].value,'=G31+ROUNDDOWN(G31*0.1,0)')
        with fitz.open(stream=pdf,filetype='pdf') as doc:
            self.assertEqual(len(doc),4)
            self.assertIn('確認用・未承認',doc[0].get_text())

    def test_summary_only_expenses_have_no_empty_detail_sheet(self):
        e=self.estimate(0);e.categories=e.categories[1:]
        xlsx,pdf,names=build_reference_outputs(e)
        self.assertEqual(names,['見積書',SUMMARY_SHEET])
        with fitz.open(stream=pdf,filetype='pdf') as doc:self.assertEqual(len(doc),2)

    def test_untrusted_text_stays_text(self):
        e=self.estimate(1);e.categories[0].items[0].name='=HYPERLINK("https://invalid.example","bad")'
        e.categories[0].items[0].spec='=1+1'
        xlsx,_,_=build_reference_outputs(e)
        w=load_workbook(BytesIO(xlsx))
        self.assertEqual(w['明細1']['B4'].data_type,'s')
        self.assertEqual(w['明細1']['C4'].data_type,'s')

    def test_local_source_reader_handles_separate_quantity_and_unit(self):
        e=self.estimate(3);_,pdf,_=build_reference_outputs(e)
        with tempfile.TemporaryDirectory() as directory:
            p=Path(directory)/'synthetic.pdf';p.write_bytes(pdf)
            q=read_reference_quote(p)
            self.assertEqual(len(q.detail),3)
            self.assertEqual(q.detail[0]['数量'],.25)
            self.assertEqual(q.detail[0]['単位'],'㎡')
            self.assertEqual(q.discount,-1000)
            self.assertEqual(q.net_total,e.net_total)
            self.assertFalse(any('別途支払い' in r['仕様'] for r in q.summary))

    def test_preserve_source_page_boundaries(self):
        e=self.estimate(3)
        e.categories[0].items[0].source_page=3
        for item in e.categories[0].items[1:]:item.source_page=4
        self.assertEqual(len(page_plan(e)),2)


if __name__=='__main__':unittest.main()
