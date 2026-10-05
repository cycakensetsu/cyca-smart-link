import tempfile
import unittest
from io import BytesIO
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from template_fill_test_v2 import (
    BRAND_NAVY,
    COMPANY,
    COMPANY_ROW,
    COVER_END_ROW,
    CUSTOMER_ROW,
    DETAIL_SHEET,
    DETAIL_PAGE_BLOCK_ROWS,
    DETAIL_PAGE_START_ROW,
    DETAIL_PAGE_SHEET_PREFIX,
    INFO_START_ROW,
    ITEM_HEADER_ROW,
    MIN_DETAIL_PAGES,
    QUOTE_SHEET,
    SUMMARY_BAND_ROW,
    SUMMARY_SHEET,
    SUMMARY_ITEM_START_ROW,
    SUMMARY_PAGE_END_ROW,
    SUMMARY_TOTAL_ROW,
    WORK_BAND_ROW,
    Category,
    Estimate,
    LineItem,
    build_single_sheet_template,
    default_test_estimate,
    estimate_from_production,
    fill_estimate_v2,
)


def _rgb(color):
    return (getattr(color, "rgb", "") or "")[-6:]


class TemplateFillV2LayoutTest(unittest.TestCase):
    def test_template_uses_landscape_layout_and_wide_columns(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            template_path = Path(temp_dir) / "template.xlsx"
            build_single_sheet_template(template_path)
            workbook = load_workbook(template_path)

        quote = workbook[QUOTE_SHEET]
        self.assertEqual(quote.page_setup.orientation, "landscape")
        self.assertEqual(quote.page_setup.fitToWidth, 1)
        self.assertEqual(quote.page_setup.fitToHeight, 0)
        self.assertGreaterEqual(quote.column_dimensions["B"].width, 34)
        self.assertGreaterEqual(quote.column_dimensions["C"].width, 30)
        self.assertGreaterEqual(quote.column_dimensions["H"].width, 32)
        breaks = [item.id for item in quote.row_breaks.brk]
        self.assertIn(COVER_END_ROW, breaks)
        self.assertIn(SUMMARY_PAGE_END_ROW, breaks)

        header = quote.cell(ITEM_HEADER_ROW, 2)
        self.assertEqual(_rgb(header.fill.fgColor), BRAND_NAVY)
        self.assertEqual(_rgb(header.font.color), "FFFFFF")

    def test_filled_output_is_one_scroll_sheet_with_a4_landscape_page_blocks(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            template_path = Path(temp_dir) / "template.xlsx"
            build_single_sheet_template(template_path)
            result = fill_estimate_v2(
                default_test_estimate(),
                template_path=template_path,
                save_to_disk=False,
            )

        self.assertTrue(result.ok, result.errors)
        workbook = load_workbook(BytesIO(result.data))
        quote = workbook[QUOTE_SHEET]
        detail = workbook[DETAIL_SHEET]

        self.assertEqual(workbook.sheetnames, [QUOTE_SHEET, DETAIL_SHEET])
        self.assertEqual(quote.page_setup.orientation, "landscape")
        self.assertEqual(quote.page_setup.fitToWidth, 1)
        self.assertEqual(quote.page_setup.fitToHeight, 0)
        self.assertEqual(quote.page_setup.pageOrder, "downThenOver")
        self.assertEqual(detail.page_setup.orientation, "landscape")
        self.assertEqual(quote.cell(CUSTOMER_ROW, 1).value, "株式会社 貞清工務店")
        self.assertEqual(quote.cell(CUSTOMER_ROW, 4).value, "御中")
        self.assertEqual(quote.cell(COMPANY_ROW, 5).value, COMPANY["name"])
        self.assertEqual(quote.cell(SUMMARY_ITEM_START_ROW, 2).value, "仮設工事")
        self.assertEqual(quote.cell(SUMMARY_ITEM_START_ROW, 7).value, "=G67")
        self.assertEqual(quote.cell(SUMMARY_TOTAL_ROW, 7).value, "=G57+G58")
        self.assertEqual(quote.cell(13, 3).value, "=G59")
        self.assertTrue(str(quote.cell(65, 7).value).startswith("=IF("))
        self.assertTrue(str(quote.cell(67, 7).value).startswith("=SUM("))
        self.assertTrue(str(quote.cell(85, 7).value).startswith("=SUMIF("))
        for page_index in range(MIN_DETAIL_PAGES):
            page_start = DETAIL_PAGE_START_ROW + page_index * DETAIL_PAGE_BLOCK_ROWS
            self.assertFalse(quote.row_dimensions[page_start].hidden)
        self.assertEqual(detail.sheet_state, "hidden")
        self.assertTrue(str(detail.cell(2, 8).value).startswith("=IF("))
        self.assertEqual(detail.cell(2, 8).alignment.horizontal, "right")

    def test_long_estimate_adds_fifth_detail_page_below_the_four_base_pages(self):
        estimate = Estimate(
            vendor="テスト株式会社",
            customer="テスト工場 御中",
            project_name="横型見積テスト",
            issue_date="2026/08/24",
            valid_days="30日",
            categories=[
                Category(
                    "塗装工事",
                    [LineItem(f"塗装明細 {index}", "仕様", 1, "式", 1000, 1000) for index in range(1, 82)],
                )
            ],
        )

        with tempfile.TemporaryDirectory() as temp_dir:
            template_path = Path(temp_dir) / "template.xlsx"
            build_single_sheet_template(template_path)
            result = fill_estimate_v2(estimate, template_path=template_path, save_to_disk=False)

        self.assertTrue(result.ok, result.errors)
        workbook = load_workbook(BytesIO(result.data))
        quote = workbook[QUOTE_SHEET]
        fifth_start = DETAIL_PAGE_START_ROW + 4 * DETAIL_PAGE_BLOCK_ROWS
        sixth_start = DETAIL_PAGE_START_ROW + 5 * DETAIL_PAGE_BLOCK_ROWS
        self.assertFalse(quote.row_dimensions[fifth_start].hidden)
        self.assertTrue(quote.row_dimensions[sixth_start].hidden)
        self.assertEqual(quote.print_area, f"'{QUOTE_SHEET}'!$A$1:$H${fifth_start + DETAIL_PAGE_BLOCK_ROWS - 1}")

    def test_production_data_replaces_all_fixed_test_values(self):
        detail = pd.DataFrame([
            {
                "見積元": "NOKフガクエンジニアリング株式会社",
                "品名": "外壁シリコン塗装",
                "仕様": "シリコン樹脂塗料",
                "数量": 120,
                "単位": "㎡",
                "見積単価": 3500,
                "見積金額": 420000,
                "備考": "",
            },
            {
                "見積元": "NOKフガクエンジニアリング株式会社",
                "品名": "仮設足場",
                "仕様": "メッシュシート共",
                "数量": 1,
                "単位": "式",
                "見積単価": 180000,
                "見積金額": 180000,
                "備考": "",
            },
        ])
        metadata = {
            "宛名": "九州テスト工場 御中",
            "工事名称": "外壁改修工事",
            "見積日": "2026/08/25",
        }
        estimate = estimate_from_production(detail, pd.DataFrame(), metadata=metadata)

        with tempfile.TemporaryDirectory() as temp_dir:
            template_path = Path(temp_dir) / "template.xlsx"
            build_single_sheet_template(template_path)
            result = fill_estimate_v2(estimate, template_path=template_path, save_to_disk=False)

        self.assertTrue(result.ok, result.errors)
        workbook = load_workbook(BytesIO(result.data))
        quote = workbook[QUOTE_SHEET]
        self.assertEqual(quote.cell(CUSTOMER_ROW, 1).value, "九州テスト工場")
        self.assertEqual(quote.cell(INFO_START_ROW, 3).value, "外壁改修工事")
        self.assertEqual(quote.cell(SUMMARY_ITEM_START_ROW, 2).value, "外壁改修工事")
        self.assertNotIn("貞清工務店", " ".join(str(cell.value or "") for row in quote.iter_rows() for cell in row))


if __name__ == "__main__":
    unittest.main()
