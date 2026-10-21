import copy
import hashlib
from io import BytesIO
from pathlib import Path
import tempfile
import unittest

from openpyxl import load_workbook

from estimate_pipeline import build_intermediate_dataframe
from estimate_review import (OriginalVault, Target, review_rows, build_draft, approve,
                             verify_approval, build_approved_estimate, target_total)


class EstimateReviewTest(unittest.TestCase):
    def setUp(self):
        self.source = {"id": hashlib.sha256(b"original").hexdigest(), "name": "a.pdf", "size": 8,
                       "tax_basis": "税抜", "tax_rate": .10, "confirmed": True, "net_total": 100000}
        self.detail, _ = build_intermediate_dataframe([
            {"品名": "屋上防水", "仕様": "仕様A", "数量": 10, "単位": "㎡", "単価": 10000, "金額": 100000,
             "見積元": "架空防水社", "工事種別": "防水工事", "__source_name": "a.pdf", "__page_number": 2}])
        self.rows = review_rows(self.detail, [self.source])
        self.meta = {"宛名": "確認用 御中", "工事名称": "複数社検証", "見積日": "2026/10/20"}

    def draft(self, rows=None, sources=None, targets=(), discount=0, meta=None, resolutions=None):
        return build_draft(self.rows, copy.deepcopy(rows or self.rows), sources or [self.source],
                           targets, discount, meta or self.meta, resolutions)

    def test_cost_and_sale_are_separate(self):
        before = copy.deepcopy(self.rows)
        d = self.draft(targets=[Target("全体", "全体", "上乗せ額", 10000)])
        self.assertEqual(d["unresolved"], [])
        self.assertEqual(d["payload"]["合計"]["仕入原価"], 100000)
        self.assertEqual(d["payload"]["合計"]["税抜売価"], 110000)
        self.assertEqual(self.rows, before)

    def test_only_measured_single_page_sources_fill_missing_page(self):
        detail=self.detail.copy()
        detail['ページ']=''
        single={**self.source,'pages':1}
        self.assertEqual(review_rows(detail,[single])[0]['ページ'],1)
        self.assertEqual(review_rows(detail,[{**single,'pages':2}])[0]['ページ'],'')
        self.assertEqual(review_rows(detail,[self.source])[0]['ページ'],'')
        detail['ページ']=2
        rows=review_rows(detail,[single])
        self.assertEqual(rows[0]['ページ'],2)
        draft=build_draft(rows,copy.deepcopy(rows),[single],metadata=self.meta)
        self.assertTrue(any(i['種類']=='原本参照' for i in draft['unresolved']))

    def test_three_scopes_and_three_target_methods(self):
        for scope,key in (("全体", "全体"), ("業者別", "架空防水社"), ("工種別", "防水工事")):
            for mode,value,expected in (("上乗せ額",10000,110000),("掛け率",1.2,120000),("粗利益率",20,125000)):
                with self.subTest(scope=scope, mode=mode):
                    d=self.draft(targets=[Target(scope,key,mode,value)])
                    self.assertEqual(d["unresolved"], [])
                    self.assertEqual(d["payload"]["合計"]["税抜売価"],expected)

    def test_discount_after_markup_and_tax(self):
        m={**self.meta,"値引き理由":"端数値引き"}
        d=self.draft(targets=[Target("全体","全体","上乗せ額",10000)], discount=-3000,meta=m)
        self.assertFalse(d["unresolved"])
        self.assertEqual(d["payload"]["合計"]["税抜売価"],107000)
        self.assertEqual(d["payload"]["合計"]["税額"],10700)
        self.assertEqual(d["payload"]["合計"]["粗利益"],7000)

    def test_discount_requires_reason(self):
        self.assertTrue(self.draft(discount=-1)["unresolved"])

    def test_positive_or_excess_discount_blocked(self):
        for v in (1,-100001):
            self.assertTrue(self.draft(discount=v)["unresolved"])

    def test_targets_reject_invalid_or_overlapping_inputs(self):
        for t in (Target("全体","全体","粗利益率",100),Target("全体","全体","掛け率",.5),
                  Target("全体","全体","上乗せ額",-1),Target("全体","全体","上乗せ額",float("inf"))):
            with self.assertRaises(ValueError):target_total(100000,t)
        d=self.draft(targets=[Target("全体","全体","上乗せ額",1),Target("工種別","防水工事","上乗せ額",1)])
        self.assertTrue(d["unresolved"])

    def test_allocation_limit_never_adds_fake_rows(self):
        d=self.draft(targets=[Target("全体","全体","上乗せ額",99999999)])
        # This unit price is uncapped, so use an explicit lump-sum fixture.
        rows=copy.deepcopy(self.rows)
        rows[0].update(数量=1,単位="式",原価単価=100000,訂正理由="原本の単位を訂正")
        d=self.draft(rows=rows,targets=[Target("全体","全体","上乗せ額",60000)])
        self.assertTrue(d["unresolved"])
        self.assertEqual(len(d["payload"]["計算明細"]),1)

    def test_included_unknown_tax_is_blocked(self):
        for basis in ("税込","不明",None):
            self.assertTrue(self.draft(sources=[{**self.source,"tax_basis":basis}])["unresolved"])

    def test_source_totals_and_tax_equation_blocked(self):
        for changes in ({"net_total":1},{"gross_total":99999},{"tax":10000,"gross_total":99999},
                        {"tax":8000,"gross_total":108000},{"confirmed":False}):
            self.assertTrue(self.draft(sources=[{**self.source,**changes}])["unresolved"])

    def test_unknown_vendor_and_trade_are_blocked(self):
        for col,val in (("工事種別","未分類"),("見積元","不明")):
            rows=copy.deepcopy(self.rows);rows[0][col]=val;rows[0]["訂正理由"]="確認"
            self.assertTrue(self.draft(rows=rows)["unresolved"])

    def test_correction_requires_reason(self):
        rows=copy.deepcopy(self.rows);rows[0]["品名"]="訂正した品名"
        self.assertTrue(self.draft(rows=rows)["unresolved"])
        rows[0]["訂正理由"]="原本2ページの文字を確認"
        self.assertFalse(self.draft(rows=rows)["unresolved"])

    def test_missing_and_nonfinite_values_cannot_be_approved(self):
        for value in (None,float("inf"),"abc"):
            rows=copy.deepcopy(self.rows);rows[0]["数量"]=value;rows[0]["訂正理由"]="確認"
            if value == float("inf"):
                with self.assertRaises(ValueError):self.draft(rows=rows)
                continue
            d=self.draft(rows=rows)
            self.assertTrue(d["unresolved"])
            with self.assertRaises(ValueError):approve(d,"検証者")

    def test_no_new_or_deleted_or_rebound_source_rows(self):
        with self.assertRaises(ValueError):self.draft(rows=self.rows*2)
        rows=copy.deepcopy(self.rows);rows[0]["原本ID"]="other"
        with self.assertRaises(ValueError):self.draft(rows=rows)

    def test_stale_and_tampered_approvals_rejected(self):
        d=self.draft();a=approve(d,"検証者");verify_approval(d,a)
        changed=self.draft(targets=[Target("全体","全体","上乗せ額",1000)])
        with self.assertRaises(ValueError):verify_approval(changed,a)
        a["actor"]="別人"
        with self.assertRaises(ValueError):verify_approval(d,a)

    def test_blank_actor_rejected(self):
        with self.assertRaises(ValueError):approve(self.draft(),"")

    def test_customer_output_contains_only_sale_prices(self):
        d=self.draft(targets=[Target("全体","全体","上乗せ額",10000)])
        result=build_approved_estimate(d,approve(d,"検証者"))
        self.assertTrue(result.ok,result.errors)
        wb=load_workbook(BytesIO(result.data))
        self.assertEqual(wb['見積書']['C13'].value, "='明細1'!G24")
        self.assertEqual(wb['見積書']['A2'].fill.fgColor.rgb[-6:], '1F3557')
        values=" ".join(str(c.value or "") for s in wb for row in s for c in row)
        for forbidden in ("原価単価","原価金額","架空防水社","a.pdf"):
            self.assertNotIn(forbidden,values)

    def test_summary_only_expense_keeps_formula(self):
        rows=copy.deepcopy(self.rows);rows[0]["集計のみ"]=True
        d=self.draft(rows=rows)
        result=build_approved_estimate(d,approve(d,"検証者"))
        self.assertTrue(result.ok,result.errors)
        wb=load_workbook(BytesIO(result.data))
        self.assertEqual(wb['工事品目集計']['G3'].value,'=ROUND(D3*F3,0)')
        self.assertEqual(wb.sheetnames, ['見積書', '工事品目集計'])

    def test_original_vault_is_immutable_and_names_cannot_escape(self):
        with tempfile.TemporaryDirectory() as temp:
            v=OriginalVault(temp,"actor")
            s=v.preserve("../../bad.pdf",b"PDF bytes")
            self.assertEqual(v.read(s["id"]),b"PDF bytes")
            self.assertFalse((Path(temp).parent/'bad.pdf').exists())
            self.assertEqual(v.preserve("again.pdf",b"PDF bytes")["id"],s["id"])
            path=v.root/'originals'/s["id"]/'original'
            path.write_bytes(b"tampered")
            with self.assertRaises(ValueError):v.read(s["id"])
            with self.assertRaises(ValueError):v.preserve("a.pdf",b"PDF bytes")

    def test_supplier_discount_is_preserved_separately_from_customer_discount(self):
        detail,_=build_intermediate_dataframe([
            dict(見積元='架空社',品名='工事',数量=10,単位='㎡',単価=10000,金額=100000,__source_name='a.pdf',__page_number=2,工事種別='防水工事'),
            dict(見積元='架空社',品名='仕入値引き',数量=1,単位='式',単価=-10000,金額=-10000,__source_name='a.pdf',__page_number=2,工事種別='防水工事')])
        rows=review_rows(detail,[self.source])
        d=build_draft(rows,rows,[{**self.source,'net_total':90000}],
                      [Target('全体','全体','上乗せ額',10000)],-1000,{**self.meta,'値引き理由':'顧客値引き'})
        self.assertFalse(d['unresolved'])
        self.assertEqual(d['payload']['計算明細'][1]['見積金額'],-10000)
        self.assertEqual(d['payload']['合計']['仕入原価'],90000)
        self.assertEqual(d['payload']['合計']['税抜売価'],99000)

    def test_reusing_same_document_is_blocked(self):
        self.assertTrue(self.draft(sources=[self.source,self.source])['unresolved'])

    def test_vault_tenant_isolation(self):
        with tempfile.TemporaryDirectory() as temp:
            a=OriginalVault(temp,'a');b=OriginalVault(temp,'b')
            s=a.preserve('a.pdf',b'only-a')
            with self.assertRaises(FileNotFoundError):b.read(s['id'])

    def test_partial_read_or_missing_document_cannot_be_approved_as_zero(self):
        source={**self.source,'read_errors':[{'page':2}]}
        self.assertTrue(any(i['種類']=='読取未完了' for i in self.draft(sources=[source])['unresolved']))
        empty={**self.source,'id':hashlib.sha256(b'missing').hexdigest(),'name':'missing.pdf','net_total':0}
        self.assertTrue(any(i['種類']=='原本明細なし' for i in self.draft(sources=[self.source,empty])['unresolved']))

    def test_source_page_must_exist_and_correction_requires_reason(self):
        rows=copy.deepcopy(self.rows);rows[0]['ページ']=3
        self.assertTrue(self.draft(rows=rows,sources=[{**self.source,'pages':2}])['unresolved'])
        rows[0]['ページ']=1;rows[0]['訂正理由']='原本1ページに掲載を確認'
        self.assertFalse(self.draft(rows=rows,sources=[{**self.source,'pages':2}])['unresolved'])


class MultiVendorReviewTest(unittest.TestCase):
    def setUp(self):
        self.sources=[];records=[]
        for i,(vendor,trade,name) in enumerate((("防水社","防水工事","ウレタン防水"),("足場社","足場工事","足場"),
                   ("塗装社","塗装工事","外壁塗装"),("板金社","板金工事","笠木"),("電気社","電気工事","配線"))):
            file=f'{i}.pdf';sid=hashlib.sha256(file.encode()).hexdigest()
            self.sources.append(dict(id=sid,name=file,size=1,tax_basis="税抜",tax_rate=.1,confirmed=True,net_total=110000))
            for item,amount in ((name,100000),("運搬費",10000)):
                records.append(dict(見積元=vendor,工事種別=trade,品名=item,仕様="同一範囲",数量=1,単位="式",単価=amount,金額=amount,__source_name=file,__page_number=2))
        details,_=build_intermediate_dataframe(records)
        self.rows=review_rows(details,self.sources)

    def test_five_trades_require_duplicate_resolutions(self):
        d=build_draft(self.rows,self.rows,self.sources)
        self.assertTrue(d['unresolved'])
        self.assertTrue(any(i['種類']=='共通費重複疑い' for i in d['unresolved']))
        answers={i['id']:'各社の搬入対象が別。原本2ページと範囲を照合。' for i in d['payload']['指摘']}
        d=build_draft(self.rows,self.rows,self.sources,resolutions=answers)
        self.assertFalse(d['unresolved'])
        self.assertEqual(d['payload']['合計']['仕入原価'],550000)

    def test_exclusion_preserves_raw_source_and_requires_reason(self):
        edited=copy.deepcopy(self.rows);edited[-1]['採用']=False
        d=build_draft(self.rows,edited,self.sources)
        self.assertTrue(any(i['種類']=='除外理由' for i in d['unresolved']))
        edited[-1]['除外理由']='他社に一括発注するため除外'
        d=build_draft(self.rows,edited,self.sources)
        self.assertEqual(len(d['payload']['原本明細']),10)
        self.assertEqual(len(d['payload']['計算明細']),9)
        self.assertEqual(d['payload']['合計']['仕入原価'],540000)

    def test_different_tax_rates_need_confirmation(self):
        sources=copy.deepcopy(self.sources);sources[-1]['tax_rate']=.08
        d=build_draft(self.rows,self.rows,sources)
        self.assertTrue(any(i['種類']=='税率混在' for i in d['unresolved']))

    def test_same_spec_sale_price_difference_needs_confirmation(self):
        edited=copy.deepcopy(self.rows)
        edited[0].update(品名='共通作業',仕様='仕様A',訂正理由='原本照合')
        edited[2].update(品名='共通作業',仕様='仕様A',訂正理由='原本照合')
        d=build_draft(self.rows,edited,self.sources,[Target('業者別','防水社','上乗せ額',10000)])
        self.assertTrue(any(i['種類']=='売価単価差' for i in d['unresolved']))


if __name__ == '__main__':
    unittest.main()
