"""CYCA enterprise reference layout: cover, 30-row work table, numbered details.

The XLSX and PDF share a page plan. Only documented source lines appear; summary
rows link to detail amounts, while original overhead lines can stay on summary.
"""
from io import BytesIO
import re
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.page import PageMargins

HEADERS = ['No', '工事品目', '仕様', '数量', '単位', '単価', '金額', '備考']
WIDTHS = [0.025, .207, .218, .049, .030, .094, .094, .283]
# Normalize the source's eight-column proportions.
WIDTHS = [x / sum(WIDTHS) for x in WIDTHS]
NAVY = '07517B'
MAX_ROWS = 26
SUMMARY_SHEET = '工事品目集計'  # Native Numbers import preserves refs to this tested name.
FONT = 'Hiragino Mincho ProN'


def page_plan(estimate):
    logical = []
    no = 0
    source_page = 0
    for index, cat in enumerate(estimate.categories):
        if cat.summary_only:
            continue
        last = None
        for item in cat.items:
            if source_page and item.source_page and source_page != item.source_page:
                logical.append({'kind':'pagebreak'})
                last = None
            if item.source_page:
                source_page = item.source_page
            section = item.section or cat.name
            if section != last:
                logical.append({'kind':'heading', 'values':['', '【'+section+'】', '', '', '', '', '', '']})
                last = section
            no += 1
            logical.append({'kind':'item', 'category':index, 'item':item,
                            'values':[no,item.name,item.spec,item.qty,item.unit,item.unit_price,item.amount,item.remark]})
    pages=[]
    while logical:
        breaks=[i for i,r in enumerate(logical) if r['kind']=='pagebreak']
        take=min(MAX_ROWS,breaks[0] if breaks else len(logical))
        if take==0:
            logical=logical[1:]
            continue
        if take < len(logical) and logical[take-1]['kind']=='heading':
            take-=1
        pages.append(logical[:take]);logical=logical[take:]
    return pages


def build_reference_outputs(estimate, draft=False):
    from template_fill_test_v2 import COMPANY
    from reportlab.pdfgen import canvas
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.lib.colors import HexColor, black, white
    from reportlab.platypus import Paragraph
    from reportlab.lib.styles import ParagraphStyle

    pages=page_plan(estimate)
    wb=Workbook();wb.remove(wb.active)
    cover=wb.create_sheet('見積書');summary=wb.create_sheet(SUMMARY_SHEET)
    thin=Side(style='thin',color='B7B7B7')
    border=Border(left=thin,right=thin,top=thin,bottom=thin)
    headerfill=PatternFill('solid',fgColor=NAVY)
    gray=PatternFill('solid',fgColor='F3F3F3')
    def set_cell(ws,row,col,value,formula=False):
        cell=ws.cell(row,col,value)
        if isinstance(value,str) and not formula:
            cell.data_type='s'
        cell.font=Font(name=FONT,size=9,color='000000')
        cell.alignment=Alignment(horizontal='right' if col in (4,6,7) else 'left',vertical='center',wrap_text=True)
        if isinstance(value,(int,float)) or formula:
            fractional=isinstance(value,float) and not value.is_integer()
            cell.number_format='#,##0.00' if col==6 and fractional else '#,##0.########' if col==4 and fractional else '#,##0'
        return cell
    def setup(ws):
        ws.sheet_view.showGridLines=False
        for i,width in enumerate(WIDTHS,1):ws.column_dimensions[chr(64+i)].width=145*width
        ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth=ws.page_setup.fitToHeight=1
        ws.sheet_properties.pageSetUpPr.fitToPage=True
        ws.page_margins=PageMargins(left=.4,right=.4,top=.4,bottom=.4,header=.1,footer=.1)
        ws.print_area='A1:H31'
        ws.oddFooter.right.text=f'{len(wb.worksheets)} / {2+len(pages)}'
        for row in range(1,32):ws.row_dimensions[row].height=15
    def table(ws,title):
        setup(ws);ws.merge_cells('A1:H1')
        set_cell(ws,1,1,title).alignment=Alignment(horizontal='center',vertical='center')
        ws.row_dimensions[1].height=21
        for col,label in enumerate(HEADERS,1):
            cell=set_cell(ws,2,col,label);cell.fill=headerfill
            cell.font=Font(name=FONT,size=9,color='FFFFFF')
            cell.alignment=Alignment(horizontal='center',vertical='center')
        for row in range(2,32):
            for col in range(1,9):
                cell=ws.cell(row,col);cell.border=border
                if row>2 and row%2==0:cell.fill=gray
        ws.merge_cells('B29:H29')
        set_cell(ws,29,2,'※工事内容明細には消費税が含まれておりません。別途支払い願います。').font=Font(name=FONT,size=7)
        ws.merge_cells('A31:B31');set_cell(ws,31,1,'合　計').alignment=Alignment(horizontal='center')

    refs={i:[] for i in range(len(estimate.categories))}
    detail_values=[]
    for page_no,page in enumerate(pages,1):
        ws=wb.create_sheet(f'明細{page_no}');table(ws,f'工 事 内 容 明 細 {page_no}')
        for row,entry in enumerate(page,3):
            for col,value in enumerate(entry['values'],1):set_cell(ws,row,col,value)
            if entry['kind']=='item':
                set_cell(ws,row,7,f'=IF(COUNT(D{row},F{row})=2,ROUND(D{row}*F{row},0),"")',True)
                refs[entry['category']].append(f"'{ws.title}'!G{row}")
        set_cell(ws,31,7,'=SUMIF(A3:A28,">0",G3:G28)',True)
        detail_values.append(page)
    table(summary,'工 事 内 容')
    # Place the familiar blank line between source categories when space allows.
    step=2 if len(estimate.categories)<=8 else 1
    summary_values=[]
    for index,cat in enumerate(estimate.categories):
        row=3+index*step
        if cat.summary_only:
            item=cat.items[0]
            values=[index+1,item.name,item.spec,item.qty,item.unit,item.unit_price,item.amount,item.remark]
            formula=f'=ROUND(D{row}*F{row},0)'
        else:
            label=cat.name if cat.name.endswith('一式') else cat.name+'一式'
            values=[index+1,label,'',1,'式',cat.subtotal,cat.subtotal,'']
            formula='=SUM('+','.join(refs[index])+')'
        for col,value in enumerate(values,1):set_cell(summary,row,col,value)
        set_cell(summary,row,7,formula,True)
        if not cat.summary_only:set_cell(summary,row,6,f'=G{row}',True)
        summary_values.append((row,values))
    subrow=max(19,3+(len(estimate.categories)-1)*step+2)
    discountrow=subrow+3
    if discountrow>28:raise ValueError('工事内容の行数が上限を超えています。工種の集約を確認してください。')
    sumrefs=','.join(f'G{row}' for row,_ in summary_values)
    set_cell(summary,subrow,1,len(estimate.categories)+1);set_cell(summary,subrow,2,'小 計')
    set_cell(summary,subrow,7,f'=SUM({sumrefs})',True)
    set_cell(summary,discountrow,2,'値 引 き');set_cell(summary,discountrow,7,estimate.discount)
    set_cell(summary,31,7,f'=G{subrow}+G{discountrow}',True)
    set_cell(summary,30,8,'税込合計')
    set_cell(summary,31,8,f'=G31+ROUNDDOWN(G31*{estimate.tax_rate},0)',True)

    setup(cover);cover.merge_cells('A2:H3')
    for col,width in zip('ABCDEFGH',[14,18,18,12,6,15,15,20]):
        cover.column_dimensions[col].width=width
    title='御　見　積　書' + ('（確認用・未承認）' if draft else '')
    c=set_cell(cover,2,1,title);c.fill=headerfill;c.font=Font(name=FONT,size=18,color='FFFFFF')
    c.alignment=Alignment(horizontal='center',vertical='center')
    cover.merge_cells('A8:C9');set_cell(cover,8,1,estimate.customer).font=Font(name=FONT,size=14)
    cover.merge_cells('A11:D11');set_cell(cover,11,1,'下記の通りお見積もり申し上げます。')
    cover.merge_cells('F5:H5');set_cell(cover,5,6,'見積日　'+estimate.issue_date)
    cover.merge_cells('F6:H6');set_cell(cover,6,6,'登録番号　'+COMPANY['reg'])
    cover.merge_cells('B16:C17');set_cell(cover,16,2,'御見積り金額').alignment=Alignment(horizontal='center')
    cover.merge_cells('D16:G17')
    grand_formula=f"='{SUMMARY_SHEET}'!H31"
    cell=set_cell(cover,16,4,grand_formula,True);cell.font=Font(name='Arial',size=22)
    cell.number_format='"¥"#,##0';cell.border=Border(bottom=Side(style='medium',color='666666'))
    cover.merge_cells('D18:G18');set_cell(cover,18,4,'消費税込').font=Font(name=FONT,size=8)
    for row,(label,value) in enumerate([('工事名称',estimate.project_name),('工事場所',''),('工事期間',''),
                        ('支払い条件','ご相談の上'),('有効期限',estimate.valid_days),('見積担当','中村 哲也')],25):
        cover.merge_cells(start_row=row,start_column=2,end_row=row,end_column=4)
        set_cell(cover,row,1,label).font=Font(name=FONT,size=8)
        set_cell(cover,row,2,value).border=Border(bottom=thin)
    for row,value in [(26,COMPANY['name']),(27,COMPANY['zip']+' '+COMPANY['address']),
                       (28,COMPANY['tel']+'　'+COMPANY['fax']),(29,COMPANY['ceo'])]:
        cover.merge_cells(start_row=row,start_column=6,end_row=row,end_column=8)
        company_cell=set_cell(cover,row,6,value)
        company_cell.font=Font(name=FONT,size=12 if row==26 else 8)
        company_cell.alignment=Alignment(horizontal='left',vertical='center',indent=1)
    for i,ws in enumerate(wb,1):ws.oddFooter.right.text=f'{i} / {len(wb.worksheets)}'
    wb.active=0;wb.calculation.calcMode='auto';wb.calculation.fullCalcOnLoad=True
    output=BytesIO();wb.save(output)

    font='HeiseiMin-W3'
    if font not in pdfmetrics.getRegisteredFontNames():pdfmetrics.registerFont(UnicodeCIDFont(font))
    pdf=BytesIO();w,h=landscape(A4);c=canvas.Canvas(pdf,pagesize=(w,h))
    c.setTitle(estimate.project_name+('（確認用・未承認）' if draft else ''))
    x=54;width=w-108;top=h-48
    def text(value,xx,yy,size=9,align='left',color=black):
        c.setFont(font,size);c.setFillColor(color)
        getattr(c,{'left':'drawString','right':'drawRightString','center':'drawCentredString'}[align])(xx,yy,str(value))
    def foot(n):
        text(('確認用・未承認　' if draft else '')+f'{n} / {2+len(pages)}',w-54,28,7,'right')
    c.setFillColor(HexColor('#'+NAVY));c.rect(x,top-24,width,24,fill=1,stroke=0)
    text(title,w/2,top-18,16,'center',white)
    text('見積日　'+estimate.issue_date,w-60,top-55,9,'right')
    text('登録番号　'+COMPANY['reg'],w-60,top-69,8,'right')
    text(estimate.customer,x+8,top-104,12)
    c.setStrokeColor(HexColor('#777777'));c.line(x,top-110,x+320,top-110)
    text('下記の通りお見積もり申し上げます。',x,top-130,9)
    text('御見積り金額',x+150,top-237,11)
    text(f'¥{estimate.grand_total:,}',x+385,top-237,22,'center')
    c.line(x+280,top-244,x+500,top-244);text('消費税込',x+500,top-260,8,'right')
    for i,(label,value) in enumerate([('工事名称',estimate.project_name),('工事場所',''),('工事期間',''),
                         ('支払い条件','ご相談の上'),('有効期限',estimate.valid_days),('見積担当','中村 哲也')]):
        yy=158-i*16;text(label,x,yy,8);text(value,x+85,yy,8);c.line(x,yy-3,x+325,yy-3)
    for yy,value,size in [(135,COMPANY['name'],18),(115,COMPANY['zip']+' '+COMPANY['address'],8),
                           (100,COMPANY['tel']+'　'+COMPANY['fax'],8),(85,COMPANY['ceo'],10)]:
        text(value,w-300,yy,size)
    foot(1);c.showPage()
    style=ParagraphStyle('source-cell',fontName=font,fontSize=8,leading=8,wordWrap='CJK')
    rowheight=16
    def render_table(title,entries,total,page):
        yy=h-72
        text(title,w/2,yy+12,10,'center')
        values_by_row=dict(entries)
        for row in range(2,32):
            if row==29:
                c.setFillColor(white);c.setStrokeColor(HexColor('#B7B7B7'))
                c.rect(x,yy-rowheight,width,rowheight,fill=1,stroke=1)
                text('※工事内容明細には消費税が含まれておりません。別途支払い願います。',x+18,yy-11,6)
                yy-=rowheight
                continue
            vals=HEADERS if row==2 else values_by_row.get(row,['']*8)
            if row==29:vals=['','※工事内容明細には消費税が含まれておりません。別途支払い願います。','','','','','','']
            if row==30 and page==2:vals=['','','','','','','','税込合計']
            if row==31:vals=['','合　計','','','','',total,estimate.grand_total if page==2 else '']
            xx=x
            for col,(val,fraction) in enumerate(zip(vals,WIDTHS),1):
                cw=width*fraction
                c.setFillColor(HexColor('#'+NAVY) if row==2 else HexColor('#F3F3F3') if row%2==0 else white)
                c.setStrokeColor(HexColor('#B7B7B7'));c.setLineWidth(.35)
                c.rect(xx,yy-rowheight,cw,rowheight,fill=1,stroke=1)
                if val != '' and val is not None:
                    if row==29:
                        if col==2:text(val,xx+3,yy-10,6)
                    elif row==2:text(val,xx+cw/2,yy-10,8,'center',white)
                    elif isinstance(val,(int,float)):
                        text(f'{val:,.0f}' if float(val).is_integer() else f'{val:,.2f}'.rstrip('0').rstrip('.'),xx+cw-3,yy-10,8,'right')
                    else:
                        para=Paragraph(escape(str(val)),style);_,ph=para.wrap(cw-6,rowheight)
                        if ph>rowheight:
                            raise ValueError('PDFの行に収まらない項目があります。仕様・品名を原本で確認してください。')
                        para.drawOn(c,xx+3,yy-rowheight/2-ph/2)
                xx+=cw
            yy-=rowheight
        foot(page);c.showPage()
    sums=list(summary_values)+[(subrow,[len(estimate.categories)+1,'小 計','','','','',estimate.subtotal,'']),
                               (discountrow,['','値 引 き','','','','',estimate.discount,''])]
    render_table('工 事 内 容',sums,estimate.net_total,2)
    for index,page in enumerate(pages,1):
        render_table(f'工 事 内 容 明 細 {index}',[(row,e['values']) for row,e in enumerate(page,3)],
                     sum(e['item'].amount for e in page if e['kind']=='item'),index+2)
    c.save()
    return output.getvalue(),pdf.getvalue(),wb.sheetnames
