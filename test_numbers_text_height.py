"""Imported Numbers rows reserve space for all wrapped source text."""
from io import BytesIO
import unittest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment
from print_layout import fit_numbers_text_rows, fit_numbers_cover_text
from template_fill_test_v2 import Estimate, Category, LineItem, fill_estimate_v2


class NumbersTextHeightTest(unittest.TestCase):
    def test_real_long_cover_title_keeps_all_text_and_page_height(self):
        title = '復元・窓仕様・室外機 出力検証20261028・提出不可'
        estimate = Estimate('', '検証用 御中', title, '2026/10/28', '30日',
            [Category('確認工事', [LineItem('確認作業', '', 1, '式', 1000, 1000)])],
            pour_layout=True, source_count=2)
        result = fill_estimate_v2(estimate, save_to_disk=False)
        self.assertTrue(result.ok, result.errors)
        sheet = load_workbook(BytesIO(result.data))['見積書']
        self.assertEqual(sheet['C25'].value, title)
        self.assertTrue(sheet['C25'].alignment.wrap_text)
        self.assertFalse(sheet['C25'].alignment.shrink_to_fit)
        self.assertEqual(sheet['C25'].font.sz, 10)
        self.assertGreater(sheet.row_dimensions[25].height, 18)
        before = sum(sheet.row_dimensions[r].height or 20 for r in range(16, 31))
        fit_numbers_cover_text(sheet)
        self.assertEqual(sum(sheet.row_dimensions[r].height or 20 for r in range(16, 31)), before)
        self.assertEqual(sheet['E25'].value, '彩架建設株式会社')

    def sheet(self):
        sheet=Workbook().active
        from print_text_height import column_width
        for index,col in enumerate('ABCDEFGH',1):
            sheet.column_dimensions[col].width=column_width(index)
        for row in range(3,23):
            sheet.row_dimensions[row].height=21
            for col in range(1,9):
                sheet.cell(row,col).font=Font(name='游ゴシック',size=9)
                sheet.cell(row,col).alignment=Alignment(wrap_text=col in (2,3,8))
        return sheet

    def test_wrapped_text_preserves_values_fonts_and_fixed_physical_page_budget(self):
        sheet=self.sheet()
        sheet['B3']='確認工事'*8
        sheet['C3']='原本仕様'*8
        sheet['D3']=12.5;sheet['F3']=1000;sheet['G3']='=ROUND(D3*F3,0)'
        sheet['B4']='諸経費'*13
        sheet['B5']='短い品名'
        sheet.row_dimensions[5].height=19
        before_values=[[sheet.cell(r,c).value for c in range(1,9)] for r in range(3,23)]
        fit_numbers_text_rows(sheet,3,22)
        self.assertGreaterEqual(sheet.row_dimensions[3].height,32)
        self.assertGreaterEqual(sheet.row_dimensions[4].height,44)
        self.assertEqual(sheet.row_dimensions[5].height,22)
        self.assertEqual([[sheet.cell(r,c).value for c in range(1,9)] for r in range(3,23)],before_values)
        self.assertAlmostEqual(sum(sheet.row_dimensions[r].height for r in range(3,23)),440)
        self.assertTrue(all(sheet.row_dimensions[r].height>=8 for r in range(3,23)))
        for cell in ('B3','C3','B4'):
            self.assertEqual(sheet[cell].font.sz,9)
            self.assertTrue(sheet[cell].alignment.wrap_text)
            self.assertFalse(sheet[cell].alignment.shrink_to_fit)

    def test_fully_occupied_page_must_repaginate_instead_of_overflowing(self):
        sheet=self.sheet()
        for row in range(3,23):sheet.cell(row,2,'確認工事')
        sheet['H8']='原本の条件\n二行目の条件\n三行目の条件'
        with self.assertRaisesRegex(ValueError,'ページの高さ'):
            fit_numbers_text_rows(sheet,3,22)
        self.assertGreaterEqual(sheet.row_dimensions[8].height,44)
        self.assertEqual(sheet['H8'].value,'原本の条件\n二行目の条件\n三行目の条件')
        self.assertEqual(sheet['H8'].font.sz,9)

    def test_customer_outputs_apply_numbers_height_to_summary_and_detail(self):
        name='確認作業'*8
        spec='原本仕様'*8
        estimate=Estimate('', '架空宛先 御中','行高の検証・提出不可','2026/10/22','30日',[
            Category('塗装工事',[LineItem(name,spec,12.5,'㎡',1000,12500)]),
            Category('原本諸経費',[LineItem(name,spec,1,'式',1000,1000)],True)],
            pour_layout=True,source_count=2)
        result=fill_estimate_v2(estimate,save_to_disk=False)
        self.assertTrue(result.ok,result.errors)
        workbook=load_workbook(BytesIO(result.data))
        summary=workbook['工事品目集計'];detail=workbook['明細1']
        self.assertEqual(summary['B4'].value,name)
        self.assertEqual(summary['C4'].value,spec)
        self.assertGreaterEqual(summary.row_dimensions[4].height,32)
        self.assertEqual(summary['G4'].value,'=ROUND(D4*F4,0)')
        self.assertEqual(detail['B4'].value,name)
        self.assertEqual(detail['C4'].value,spec)
        self.assertGreaterEqual(detail.row_dimensions[4].height,32)
        self.assertEqual(detail['D4'].value,12.5)
        self.assertEqual(detail['G4'].value,'=ROUND(D4*F4,0)')
        self.assertEqual(len(workbook.sheetnames),3)
        for sheet in workbook:
            self.assertEqual(sheet.page_setup.fitToWidth,1)
            self.assertEqual(sheet.page_setup.fitToHeight,1)
