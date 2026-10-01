import unittest
from io import BytesIO

from openpyxl import load_workbook

from estimate_pipeline import (
    apply_profit,
    build_cost_basis_dataframe,
    build_intermediate_dataframe,
)
from horizontal_estimate_flow import build_horizontal_estimate


class HorizontalEstimateFlowTest(unittest.TestCase):
    def setUp(self):
        self.details, _ = build_intermediate_dataframe([
            {"No": 1, "見積元": "新規塗装株式会社", "品名": "外壁塗装", "数量": "2式", "単価": 100000, "金額": 200000},
            {"No": 1, "見積元": "新規設備株式会社", "品名": "設備移設", "数量": "1式", "単価": 90000, "金額": 90000},
        ])
        self.costs, _ = build_cost_basis_dataframe({}, self.details)
        self.metadata = {"宛名": "今回の依頼主 御中", "工事名称": "今回の新規工事"}

    def test_new_upload_without_markup_replaces_demo_values(self):
        quoted = apply_profit(self.costs, "上乗せしない（原価そのまま）")
        result = build_horizontal_estimate(self.details, self.costs, quoted, self.metadata)
        self.assertTrue(result.ok, result.errors)
        self.assertEqual(result.summary["小計"], 290000)
        workbook = load_workbook(BytesIO(result.data))
        self.assertEqual(workbook.sheetnames, ["見積書", "明細データ"])
        values = " ".join(str(cell.value or "") for row in workbook["見積書"] for cell in row)
        self.assertIn("今回の新規工事", values)
        self.assertIn("新規塗装株式会社", values)
        self.assertIn("新規設備株式会社", values)
        self.assertNotIn("貞清工務店", values)

    def test_standard_markup_reaches_horizontal_workbook(self):
        quoted = apply_profit(self.costs, "固定金額（円）を全体に割り振る", 10000)
        result = build_horizontal_estimate(self.details, self.costs, quoted, self.metadata)
        self.assertTrue(result.ok, result.errors)
        self.assertEqual(result.summary["小計"], int(quoted["見積金額"].sum()))

    def test_empty_details_are_not_replaced_by_demo_data(self):
        with self.assertRaisesRegex(ValueError, "明細"):
            build_horizontal_estimate(self.details.iloc[0:0], self.costs, self.costs, self.metadata)


if __name__ == "__main__":
    unittest.main()
