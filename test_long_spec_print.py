from io import BytesIO
import re
import unittest
import fitz
from openpyxl import load_workbook
from template_fill_test_v2 import Estimate,Category,LineItem,fill_estimate_v2
from pour_layout import plan
from print_text_height import entry_height,grid_fits


class LongSpecPrintTest(unittest.TestCase):
    def estimate(self,count=10):
        spec=('採寸W:520mm 採寸H:1425mm 出来寸W:510mm 出来寸H:1415mm '+
              '化粧カバー・樹脂カバー・気密テープ30mm/50mm/75mm。')*4+'原本寸法W493×H1399mm 障子部W423×H1301mm'
        return Estimate('', '架空の検証先 御中','長文仕様検証・提出不可','2026/10/27','30日',[
            Category('サッシ工事',[LineItem('検証窓'+str(i),spec if i%3==0 else '短い仕様',2,'窓',100,200,row_id='row'+str(i))
                                  for i in range(count)])],pour_layout=True,source_count=1)

    def test_long_spec_reflows_shared_pages_without_losing_text_rows_money_or_font_size(self):
        estimate=self.estimate(18)
        layout=plan(estimate)
        self.assertTrue(all(grid_fits([entry_height(r) for r in p]) for p in layout['pages']))
        self.assertTrue(all(p[-1]['kind'] not in ('category','section') for p in layout['pages']))
        self.assertEqual([r['item'].row_id for p in layout['pages'] for r in p if r['kind']=='item'],
                         ['row'+str(i) for i in range(18)])
        result=fill_estimate_v2(estimate,save_to_disk=False)
        self.assertTrue(result.ok,result.errors)
        workbook=load_workbook(BytesIO(result.data))
        with fitz.open(stream=result.pdf_data,filetype='pdf') as pdf:
            self.assertEqual(len(pdf),len(workbook.worksheets))
            text=''.join(p.get_text() for p in pdf)
            stripped=re.sub(r'\s','',text)
            self.assertEqual(stripped.count(re.sub(r'\s','',estimate.categories[0].items[0].spec)),6)
            for i in range(18):self.assertIn('検証窓'+str(i),stripped)
            for page in pdf:
                for block in page.get_text('dict')['blocks']:
                    for line in block.get('lines',[]):
                        for span in line['spans']:
                            self.assertGreaterEqual(span['bbox'][1],0)
                            self.assertLessEqual(span['bbox'][3],page.rect.height)
        printed=[]
        for sheet in workbook:
            if not sheet.title.startswith('明細'):continue
            self.assertLessEqual(sum(sheet.row_dimensions[r].height or 21 for r in range(3,23)),420.001)
            for row in range(3,23):
                if isinstance(sheet.cell(row,1).value,int):
                    printed.append(sheet.cell(row,2).value)
                    self.assertEqual(sheet.cell(row,7).value,f'=ROUND(D{row}*F{row},0)')
                    self.assertEqual(sheet.cell(row,3).font.sz,9)
        self.assertEqual(printed,['検証窓'+str(i) for i in range(18)])

    def test_long_summary_expenses_gain_pages_and_keep_cross_sheet_totals(self):
        estimate=self.estimate(1)
        estimate.source_count=2
        estimate.categories += [Category('費用'+str(i),[LineItem('原本に記載した養生・安全対策・福利費'*3,'条件を省略しない',1,'式',100,100)],True)
                                for i in range(16)]
        layout=plan(estimate)
        self.assertGreater(layout['summary_pages'],1)
        result=fill_estimate_v2(estimate,save_to_disk=False)
        self.assertTrue(result.ok,result.errors)
        workbook=load_workbook(BytesIO(result.data))
        self.assertIn('工事品目集計2',workbook.sheetnames)
        self.assertIn("'工事品目集計2'!G3:G18",workbook['工事品目集計']['G22'].value)

    def test_one_item_beyond_one_physical_page_stops_without_truncation_or_partial_output(self):
        estimate=self.estimate(1)
        estimate.categories[0].items[0].spec='省略不可の原本仕様'*1000
        result=fill_estimate_v2(estimate,save_to_disk=False)
        self.assertFalse(result.ok)
        self.assertIsNone(result.data)
        self.assertIsNone(result.pdf_data)
        self.assertTrue(any('1ページに収まりません' in e for e in result.errors))

    def test_copied_summary_and_late_detail_sheets_keep_both_footer_fields_and_total_label(self):
        estimate=self.estimate(180)
        for item in estimate.categories[0].items:item.spec='短い仕様'
        estimate.source_count=2
        estimate.categories += [Category('費用'+str(i),[LineItem('検証費用'+str(i),'',1,'式',100,100)],True)
                                for i in range(16)]
        result=fill_estimate_v2(estimate,save_to_disk=False)
        self.assertTrue(result.ok,result.errors)
        workbook=load_workbook(BytesIO(result.data))
        self.assertIn('工事品目集計2',workbook.sheetnames)
        self.assertIn('明細9',workbook.sheetnames)
        for i,sheet in enumerate(workbook,1):
            self.assertIn(estimate.project_name,sheet.oddFooter.left.text)
            self.assertIsNone(sheet.oddFooter.left.size)
            self.assertEqual(sheet.oddFooter.right.text,f'{i} / {len(workbook.worksheets)}')
            self.assertEqual(sheet.oddFooter.right.size,8)
            if sheet.title.startswith('工事品目集計'):
                self.assertEqual(sheet['E22'].value,'全体小計')
        with fitz.open(stream=result.pdf_data,filetype='pdf') as pdf:
            for page in (pdf[1],pdf[2]):
                self.assertIn('合計欄は、まとめ全ページを合わせた金額です。',page.get_text())


if __name__=='__main__':unittest.main()
