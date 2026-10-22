"""Printed context must stop at source/part boundaries without leaking issuers."""
from io import BytesIO
import unittest
import fitz
from openpyxl import load_workbook
from template_fill_test_v2 import Estimate, Category, LineItem, fill_estimate_v2
from estimate_review import estimate_for_output, _customer_source_project
from pour_layout import plan


def example(*, long=False):
    count=18 if long else 2
    items=[LineItem(f'車庫作業{n+1}','原本仕様',1,'式',1000,1000,
                    section='車庫',source_id='private-source-a',row_id=f'a{n}') for n in range(count)]
    items += [LineItem(f'別原本の作業{n+1}','原本仕様',1,'式',2000,2000,
                      source_id='private-source-b',row_id=f'b{n}') for n in range(4)]
    return Estimate('非公開施工社','架空宛先 御中','境界検証・提出不可','2026/10/21','30日',
                    [Category('塗装工事',items)],pour_layout=True,source_count=2)


class SourceContextLayoutTest(unittest.TestCase):
    def test_unnamed_new_source_resets_prior_garage_scope_in_xlsx_and_pdf(self):
        estimate=example()
        layout=plan(estimate)
        rows=layout['pages'][0]
        first_new=next(i for i,r in enumerate(rows) if r['kind']=='item' and r['item'].row_id=='b0')
        self.assertEqual(rows[first_new-1],{'kind':'section','label':'塗装工事'})
        self.assertEqual([r['no'] for page in layout['pages'] for r in page if r['kind']=='item'],list(range(1,7)))
        result=fill_estimate_v2(estimate,save_to_disk=False)
        self.assertTrue(result.ok,result.errors)
        workbook=load_workbook(BytesIO(result.data))
        values=[workbook['明細1'].cell(n,2).value for n in range(3,23)]
        self.assertLess(values.index('【車庫】'),values.index('【塗装工事】'))
        self.assertLess(values.index('【塗装工事】'),values.index('別原本の作業1'))
        all_values=' '.join(str(c.value or '') for sheet in workbook for row in sheet for c in row)
        with fitz.open(stream=result.pdf_data,filetype='pdf') as pdf:
            text=''.join(p.get_text() for p in pdf)
            self.assertLess(text.index('【車庫】'),text.index('【塗装工事】'))
            self.assertLess(text.index('【塗装工事】'),text.index('別原本の作業1'))
            for private in ('private-source-a','private-source-b','非公開施工社'):
                self.assertNotIn(private,text)
                self.assertNotIn(private,all_values)

    def test_page_continuation_keeps_the_new_source_context_and_twenty_row_budget(self):
        estimate=example(long=True)
        estimate.categories[0].items[-1].source_project='外壁補修工事'
        pages=plan(estimate)['pages']
        self.assertTrue(all(len(page)<=20 for page in pages))
        self.assertTrue(all(page[-1]['kind'] not in ('category','section') for page in pages))
        self.assertEqual(sum(r['kind']=='item' for page in pages for r in page),22)
        flat=[r for page in pages for r in page]
        index=next(n for n,r in enumerate(flat) if r['kind']=='item' and r['item'].row_id=='b0')
        self.assertEqual(flat[index-1]['label'],'塗装工事')
        self.assertFalse(any('車庫' in page[0].get('label','') for page in pages if any(r['kind']=='item' and r['item'].row_id.startswith('b') for r in page) and not any(r['kind']=='item' and r['item'].row_id.startswith('a') for r in page)))
        result=fill_estimate_v2(estimate,save_to_disk=False)
        self.assertTrue(result.ok,result.errors)

    def test_explicit_source_title_is_optional_and_company_or_filename_titles_are_not_printed(self):
        safe={'id':'a','project_name':'外壁補修工事','vendor_name':'非公開施工社'}
        for title in ('非公開施工社の外壁補修工事','非公開 施工社の外壁補修工事','株式会社非公開の外壁補修工事','private-source.pdf','工事'*30):
            with self.subTest(title=title):
                self.assertEqual(_customer_source_project(dict(safe,project_name=title),[]),'')
        self.assertEqual(_customer_source_project(safe,[]),'外壁補修工事')
        row={'品名':'確認作業','数量':1,'単位':'式','見積単価':1000,'見積金額':1000,'部位':'',
             '行ID':'r1','原本ID':'a','工事種別':'塗装工事','見積元':'非公開施工社'}
        draft={'payload':{'計算明細':[row],'原本':[safe], '表題':{'宛名':'架空宛先','工事名称':'検証'},'値引き':0}}
        estimate=estimate_for_output(draft)
        item=estimate.categories[0].items[0]
        self.assertEqual(item.source_id,'a')
        self.assertEqual(item.source_project,'外壁補修工事')
        # Compact and long layouts use the same context rules.
        self.assertEqual(plan(estimate)['pages'][0][0],{'kind':'section','label':'外壁補修工事'})

    def test_same_named_part_in_different_original_gets_a_new_heading(self):
        estimate=example()
        for item in estimate.categories[0].items:item.section='屋上'
        rows=plan(estimate)['pages'][0]
        self.assertEqual([r['label'] for r in rows if r['kind']=='section'],['屋上','屋上'])
        # An explicit part followed by an unnamed part in the same original
        # also closes the previous heading, rather than inheriting its scope.
        estimate.categories[0].items[1].section=''
        rows=plan(estimate)['pages'][0]
        first_reset=next(i for i,r in enumerate(rows) if r['kind']=='item' and r['item'].row_id=='a1')
        self.assertEqual(rows[first_reset-1]['label'],'塗装工事')

    def test_safe_distinct_projects_disambiguate_the_same_named_part(self):
        estimate=example()
        for item in estimate.categories[0].items:
            item.section='屋上'
            item.source_project='東棟補修' if item.source_id=='private-source-a' else '西棟補修'
        rows=plan(estimate)['pages'][0]
        self.assertEqual([r['label'] for r in rows if r['kind']=='section'],['東棟補修／屋上','西棟補修／屋上'])
        result=fill_estimate_v2(estimate,save_to_disk=False)
        self.assertTrue(result.ok,result.errors)
